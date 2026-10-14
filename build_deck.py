import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Modern High-Tech Deep Slate
    C_BG = RGBColor(11, 15, 25)          # #0B0F19 Deep Night
    C_CARD_BG = RGBColor(22, 32, 50)     # #162032 Card Surface
    C_CARD_BORDER = RGBColor(38, 54, 82) # #263652 Card Border
    C_CARD_HOVER = RGBColor(30, 42, 66)  # Accent card
    
    C_INDIGO = RGBColor(99, 102, 241)    # #6366F1 Brand Accent
    C_CYAN = RGBColor(6, 182, 212)       # #06B6D4 Secondary Accent
    C_EMERALD = RGBColor(16, 185, 129)   # #10B981 Success
    C_AMBER = RGBColor(245, 158, 11)     # #F59E0B Warning / Notice
    C_ROSE = RGBColor(244, 63, 94)       # #F43F5E Coral / Alert
    C_PURPLE = RGBColor(168, 85, 247)    # #A855F7 Violet

    C_TEXT_WHITE = RGBColor(248, 250, 252) # #F8FAFC
    C_TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
    C_TEXT_DIM = RGBColor(100, 116, 139)   # #64748B
    C_TEXT_HIGHLIGHT = RGBColor(226, 232, 240) # #E2E8F0

    FONT_TITLE = "Microsoft YaHei"
    FONT_BODY = "Microsoft YaHei"

    def set_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background() # no line
        return bg

    def add_header(slide, badge_text, title_text, subtitle_text, badge_color=C_INDIGO):
        # Badge
        badge_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.48), Inches(2.6), Inches(0.36))
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = badge_color
        badge_box.line.fill.background()
        tf_b = badge_box.text_frame
        tf_b.word_wrap = True
        tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_b = tf_b.paragraphs[0]
        p_b.text = badge_text
        p_b.font.name = FONT_TITLE
        p_b.font.size = Pt(10.5)
        p_b.font.bold = True
        p_b.font.color.rgb = C_TEXT_WHITE
        p_b.alignment = PP_ALIGN.CENTER

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.86), Inches(11.7), Inches(0.6))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_TITLE
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = C_TEXT_WHITE

        # Subtitle
        s_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.44), Inches(11.7), Inches(0.45))
        tf_s = s_box.text_frame
        tf_s.word_wrap = True
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(11.5)
        p_s.font.color.rgb = C_TEXT_MUTED

    def add_card(slide, left, top, width, height, title="", title_color=C_TEXT_WHITE, bg_color=C_CARD_BG, border_color=C_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.25)

        if title:
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = FONT_TITLE
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = title_color
            p.space_after = Pt(10)
        return card, tf

    # =========================================================================
    # SLIDE 1: COVER
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)

    # Decorative top bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_INDIGO
    bar.line.fill.background()

    # Left Column: Text & Information
    badge1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(3.2), Inches(0.38))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = C_INDIGO
    badge1.line.fill.background()
    tf = badge1.text_frame
    p = tf.paragraphs[0]
    p.text = "⚡ 高性能协议引擎 · 架构深度解析"
    p.font.name = FONT_TITLE
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Main Title
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(7.5), Inches(1.8))
    tf = t_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "KiroX"
    p.font.name = FONT_TITLE
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_WHITE

    p2 = tf.add_paragraph()
    p2.text = "新一代全协议 Kiro 注册与资产管理系统"
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = C_CYAN
    p2.space_before = Pt(8)

    # Subtitle
    sub_box = s1.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(7.5), Inches(1.1))
    tf = sub_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "基于 Wails v2 + Go + 真实 TLS 指纹模拟的高性能跨平台桌面注册工具，打造免浏览器渲染、毫秒级执行、多邮箱生态与智能住宅代理中继的专业闭环。"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.color.rgb = C_TEXT_MUTED

    # Highlights Row (4 chips)
    chips = [
        ("⚡ 纯协议零渲染", C_INDIGO),
        ("📬 5大邮箱生态", C_CYAN),
        ("🛡️ TLS 指纹对抗", C_EMERALD),
        ("🌐 住宅代理智能调度", C_AMBER)
    ]
    for i, (ctext, ccol) in enumerate(chips):
        cx = Inches(1.0 + i * 1.85)
        cy = Inches(4.85)
        cshape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(1.75), Inches(0.42))
        cshape.fill.solid()
        cshape.fill.fore_color.rgb = C_CARD_BG
        cshape.line.color.rgb = ccol
        cshape.line.width = Pt(1)
        ctf = cshape.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cp = ctf.paragraphs[0]
        cp.text = ctext
        cp.font.name = FONT_TITLE
        cp.font.size = Pt(10.5)
        cp.font.bold = True
        cp.font.color.rgb = ccol
        cp.alignment = PP_ALIGN.CENTER

    # Footer Metadata
    meta_box = s1.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(7.5), Inches(0.8))
    tf = meta_box.text_frame
    p = tf.paragraphs[0]
    p.text = "作者：1in (@huey1in)   |   版本：v1.0.4   |   协议：Apache 2.0 License"
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = C_TEXT_DIM

    # Right Column: Visual Box with Logo
    logo_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(1.3), Inches(3.6), Inches(4.9))
    logo_card.fill.solid()
    logo_card.fill.fore_color.rgb = C_CARD_BG
    logo_card.line.color.rgb = C_CARD_BORDER
    logo_card.line.width = Pt(1.5)

    # Add app icon
    icon_path = "frontend/assets/appicon.png"
    if os.path.exists(icon_path):
        s1.shapes.add_picture(icon_path, Inches(9.85), Inches(1.7), width=Inches(1.7), height=Inches(1.7))

    card_tf = logo_card.text_frame
    card_tf.word_wrap = True
    card_tf.margin_top = Inches(2.4)
    cp1 = card_tf.paragraphs[0]
    cp1.text = "KiroX Desktop"
    cp1.font.name = FONT_TITLE
    cp1.font.size = Pt(18)
    cp1.font.bold = True
    cp1.font.color.rgb = C_TEXT_WHITE
    cp1.alignment = PP_ALIGN.CENTER

    cp2 = card_tf.add_paragraph()
    cp2.text = "Cross-Platform Native Engine"
    cp2.font.name = FONT_BODY
    cp2.font.size = Pt(11)
    cp2.font.color.rgb = C_CYAN
    cp2.alignment = PP_ALIGN.CENTER
    cp2.space_before = Pt(4)

    cp3 = card_tf.add_paragraph()
    cp3.text = "Windows · macOS · Linux\nWails v2.11.0 · Go 1.24.4"
    cp3.font.name = FONT_BODY
    cp3.font.size = Pt(11)
    cp3.font.color.rgb = C_TEXT_MUTED
    cp3.alignment = PP_ALIGN.CENTER
    cp3.space_before = Pt(12)

    # =========================================================================
    # SLIDE 2: BACKGROUND & PAIN POINTS
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "背景与痛点 · PAIN POINTS", "为什么传统自动化难以胜任 Kiro 自动化注册？", "剖析 AI 编码工具与云厂商鉴权生态下，传统方案面临的严峻挑战与瓶颈")

    cols_data = [
        {
            "badge": "🛑 痛点一",
            "title": "传统 UI 自动化资源沉重易崩溃",
            "accent": C_ROSE,
            "items": [
                "Chromium 内存黑洞：单无头实例消耗 300MB~800MB 内存，单机并发超过 5 个即濒临死机崩溃。",
                "启动与执行耗时极长：需拉起完整 DOM 树与资源文件渲染，单个注册周期耗时长达 2~5 分钟。",
                "维护脆弱度极高：前端界面 DOM 样式与 Class 名稍有迭代，自动化定位选择器立即大面积失效。"
            ]
        },
        {
            "badge": "🎯 痛点二",
            "title": "反爬指纹探测与风控拦截极其严苛",
            "accent": C_AMBER,
            "items": [
                "浏览器环境特征裸露：navigator.webdriver、CDP 调试端口及渲染管线被云端风控系统秒级识别捕获。",
                "网络层缺乏指纹定制：标准 Go/Python 请求缺少 Chrome 握手套件特征，JA3 指纹直接触发 WAF 阻断。",
                "高频遭遇 OTP 400 封锁：使用脏 IP 或未伪装指纹极易被 AWS 邮箱验证码服务直接判定为恶意请求拒绝。"
            ]
        },
        {
            "badge": "🧩 痛点三",
            "title": "鉴权链路极其复杂，邮箱生态碎片化",
            "accent": C_CYAN,
            "items": [
                "全流程长达 15+ 环节：从 OIDC 设备码、邮箱双向验证，到 JWE 非对称密文加密及 SSO 令牌置换。",
                "收件协议规范碎片：企业常用 Outlook、商业临时邮箱、自建私有邮件服务缺乏统一高效的拉取机制。",
                "缺乏容错与状态追踪：传统脚本遇到网络抖动经常造成「密码已设置但令牌丢失」，白白浪费邮箱资源。"
            ]
        }
    ]

    for i, col in enumerate(cols_data):
        cx = Inches(0.8 + i * 3.98)
        cy = Inches(2.05)
        cw = Inches(3.78)
        ch = Inches(4.8)
        c, ctf = add_card(s2, cx, cy, cw, ch, col["title"], col["accent"])
        
        # Add badge
        b_p = ctf.paragraphs[0]
        b_p.text = col["badge"] + "  " + col["title"]
        b_p.font.size = Pt(13)
        b_p.font.bold = True
        b_p.font.color.rgb = col["accent"]
        b_p.space_after = Pt(14)

        for item in col["items"]:
            p = ctf.add_paragraph()
            parts = item.split("：", 1)
            p.space_after = Pt(10)
            if len(parts) == 2:
                run1 = p.add_run()
                run1.text = "• " + parts[0] + "："
                run1.font.name = FONT_TITLE
                run1.font.size = Pt(11)
                run1.font.bold = True
                run1.font.color.rgb = C_TEXT_WHITE

                run2 = p.add_run()
                run2.text = parts[1]
                run2.font.name = FONT_BODY
                run2.font.size = Pt(10.5)
                run2.font.color.rgb = C_TEXT_MUTED
            else:
                run = p.add_run()
                run.text = "• " + item
                run.font.name = FONT_BODY
                run.font.size = Pt(10.5)
                run.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 3: CORE VALUE & POSITIONING
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "核心价值 · VALUE PROPOSITION", "KiroX 破局之道：纯协议级极速注册引擎", "告别繁重浏览器模拟，通过深度协议逆向与工业级工程设计实现质的飞跃")

    quads = [
        {
            "badge": "⚡ 极致性能与低开销",
            "title": "纯协议级网络 I/O 执行",
            "color": C_INDIGO,
            "desc": "抛弃庞大的无头渲染引擎，全链路采用原生 HTTP/TLS 请求完成交互。单机可流畅并发数十个 Worker，CPU 占用极低，单实例内存仅需几十 MB，注册全流程压缩至数十秒内。"
        },
        {
            "badge": "🛡️ 深度网络与指纹抗封",
            "title": "真实 TLS / JA3 / 环境伪装",
            "color": C_CYAN,
            "desc": "基于定制 tls-client 引擎，完美还原现代 Chrome 的 TLS 握手套件、扩展顺序与 HTTP/2 首部流；内置高仿真 performance.timing 与屏幕参数动态生成模型，从容绕过风控检测。"
        },
        {
            "badge": "📬 统一全渠道邮箱矩阵",
            "title": "主流商业与私有邮件系统通收",
            "color": C_EMERALD,
            "desc": "深度集成 Outlook 邮箱池 (IMAP/Graph 双引擎)、iCloud 消息列表极速拉取、MoeMail 临时多域名池、自部署 Cloud-Mail 自动同步以及 MailNest 商业邮箱，提供高度解耦的统一收件接口。"
        },
        {
            "badge": "🔄 闭环健壮状态机",
            "title": "原子化状态与账号实时探活",
            "color": C_AMBER,
            "desc": "密码设置成功即时原子标记，网络异常自动执行三次幂等重试，杜绝邮箱无谓损耗；注册成功自动调用 Verify 探测账号存活 (alive) 与 CodeWhisperer 信用额度 (creditLimit/Used)。"
        }
    ]

    for i, q in enumerate(quads):
        row = i // 2
        col = i % 2
        qx = Inches(0.8 + col * 5.95)
        qy = Inches(2.05 + row * 2.45)
        qw = Inches(5.75)
        qh = Inches(2.25)

        c, ctf = add_card(s3, qx, qy, qw, qh, q["badge"] + "  |  " + q["title"], q["color"])
        p0 = ctf.paragraphs[0]
        p0.font.size = Pt(13.5)
        p0.font.bold = True
        p0.space_after = Pt(8)

        p = ctf.add_paragraph()
        p.text = q["desc"]
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "系统架构 · SYSTEM ARCHITECTURE", "分层解耦的工业级跨平台架构设计", "以 Go 语言构建稳健并发核心，以 Wails v2 赋能现代化敏捷交互界面")

    layers = [
        {
            "name": "🖥️ 展现与交互层 (Presentation Layer)",
            "tech": "Wails v2.11.0 + 原生 HTML5 / CSS3 / ES6 + D3.js 可视化",
            "accent": C_INDIGO,
            "details": "• 概览仪表盘、实时脱敏日志滚屏、邮箱池管理、IP 代理池测试、高级配置与版本更新\n• D3.js 实时渲染浏览器指纹偏移曲线 | 中/英/日多语言无缝切换 | 深浅自适应主题与完成提示音"
        },
        {
            "name": "🔗 核心应用桥接层 (Application Bridge)",
            "tech": "app.go Wails 绑定方法 + 事件总线 (Event Bus)",
            "accent": C_CYAN,
            "details": "• 暴露高内聚 Go 方法至前端 window.go.main.App，实现安全双向 RPC 调用\n• 异步轮询与事件推送机制，毫秒级同步任务执行进度、存活状态及敏感后端 URL 自动脱敏"
        },
        {
            "name": "⚙️ 业务领域服务层 (Domain Services)",
            "tech": "内部核心包：core / task / email / browser / proxy / storage",
            "accent": C_EMERALD,
            "details": "• task 模块：并发 Worker Pool 控制、串行等待调度、Context 优雅中断\n• core 模块：15 步注册流水线编排、重试容灾与存活校验\n• email & browser：全协议邮箱提取抽象矩阵 + 浏览器上下文特征与指纹缓存构建器"
        },
        {
            "name": "🔌 网络与安全基建层 (Infrastructure & Security)",
            "tech": "bogdanfinn/tls-client + JWE (RSA-OAEP-256 / AES-256-GCM)",
            "accent": C_AMBER,
            "details": "• 基于定制 TLS 客户端实现底层 JA3/JA4 指纹模拟与 HTTP/2 特征对齐\n• JWE 与 XXTEA 纯 Go 原生加密实现，支持 SOCKS5 / HTTP / HTTPS 代理隧道与链路延迟诊断"
        }
    ]

    for i, lyr in enumerate(layers):
        ly = Inches(2.05 + i * 1.22)
        c, ctf = add_card(s4, Inches(0.8), ly, Inches(11.7), Inches(1.12), "", C_TEXT_WHITE, C_CARD_BG, C_CARD_BORDER)
        
        p0 = ctf.paragraphs[0]
        run_name = p0.add_run()
        run_name.text = lyr["name"] + "   "
        run_name.font.name = FONT_TITLE
        run_name.font.size = Pt(12.5)
        run_name.font.bold = True
        run_name.font.color.rgb = lyr["accent"]

        run_tech = p0.add_run()
        run_tech.text = "[ " + lyr["tech"] + " ]"
        run_tech.font.name = FONT_BODY
        run_tech.font.size = Pt(11)
        run_tech.font.color.rgb = C_TEXT_MUTED
        p0.space_after = Pt(4)

        p1 = ctf.add_paragraph()
        p1.text = lyr["details"]
        p1.font.name = FONT_BODY
        p1.font.size = Pt(10)
        p1.font.color.rgb = C_TEXT_HIGHLIGHT

    # =========================================================================
    # SLIDE 5: 15-STEP PROTOCOL WORKFLOW
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "业务流转 · PROTOCOL PIPELINE", "15 步纯协议注册流水线全景拆解", "精准复刻 AWS Builder ID 认证注册与 Kiro 令牌置换的完整状态机流转")

    phases = [
        {
            "title": "阶段一：初始化与会话握手",
            "accent": C_INDIGO,
            "steps": [
                ("01 OIDC 注册", "客户端注册并换取动态 Client ID 与密钥"),
                ("02 设备注册", "申请 Device Code 及专用用户验证 URI"),
                ("03 邮箱分配", "从配置邮箱源智能调度可用邮箱实体"),
                ("04 门户交互", "预热加载远程 AppJS 配置并初始化指纹"),
                ("05 流程初始化", "初始化 AWS 注册工作流会话与上下文参数")
            ]
        },
        {
            "title": "阶段二：验证码与安全加密",
            "accent": C_CYAN,
            "steps": [
                ("06 提交邮箱", "提交账号并根据状态分流至 Signup 分支"),
                ("07~08 Profile", "用户身份配置初始化并启动注册环境"),
                ("09~10 验证码", "请求 SendOTP，多协议轮询并提取验证码"),
                ("11 创建身份", "提交有效验证码确认建立云端身份标识"),
                ("12 设置密码", "JWE 密文安全提交密码 (置位 passwordSet)")
            ]
        },
        {
            "title": "阶段三：凭证交换与存活探活",
            "accent": C_EMERALD,
            "steps": [
                ("12.8 SSO 准备", "启动 SSO 工作流验证，打通登录前序会话"),
                ("13 SSO Token", "获取 AWS SSO Access/Refresh Token (3次重试)"),
                ("14 Kiro 授权", "发起 Kiro 客户端授权，置换临时 Auth Code"),
                ("15 交换令牌", "置换最终 Kiro Refresh Token 与 ClientSecret"),
                ("校验 账号存活", "探测账号存活状态与 CodeWhisperer 信用额度")
            ]
        }
    ]

    for i, ph in enumerate(phases):
        px = Inches(0.8 + i * 3.98)
        py = Inches(2.05)
        pw = Inches(3.78)
        ph_height = Inches(4.8)

        c, ctf = add_card(s5, px, py, pw, ph_height, ph["title"], ph["accent"])
        p0 = ctf.paragraphs[0]
        p0.font.size = Pt(13)
        p0.space_after = Pt(12)

        for s_idx, (s_name, s_desc) in enumerate(ph["steps"]):
            p = ctf.add_paragraph()
            p.space_after = Pt(9)

            r1 = p.add_run()
            r1.text = s_name + "\n"
            r1.font.name = FONT_TITLE
            r1.font.size = Pt(11)
            r1.font.bold = True
            r1.font.color.rgb = ph["accent"]

            r2 = p.add_run()
            r2.text = s_desc
            r2.font.name = FONT_BODY
            r2.font.size = Pt(10)
            r2.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 6: ANTI-BOT & FINGERPRINTING
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "反爬攻防 · STEALTH ENGINEERING", "多维指纹模拟：从传输层到应用层的立体防御", "以工业级仿真标准对抗云厂商风控，保障高并发下的高通过率")

    anti_cards = [
        {
            "title": "🌐 传输层：TLS / JA3 / HTTP2 指纹仿真",
            "accent": C_INDIGO,
            "bullets": [
                "**真实 Chrome TLS 握手**：基于定制 bogdanfinn/tls-client，精准模拟真实 Chrome 浏览器的 Client Hello 特征。",
                "**密码套件与扩展对齐**：完整还原 Cipher Suites、Supported Versions、Key Share 与 Elliptic Curves 顺序，杜绝特征异常。",
                "**HTTP/2 协议层规整**：对齐 SETTINGS 帧参数与伪首部 (:method, :authority, :scheme, :path) 顺序，消除 WAF 识别漏洞。"
            ]
        },
        {
            "title": "🎭 应用层：全维度浏览器上下文构建",
            "accent": C_CYAN,
            "bullets": [
                "**时延模型高斯分布**：GenPerfTiming 算法动态模拟 DNS 查询、TCP 连接、DOM 解析等真实网络耗时，击穿静态脚本检测。",
                "**屏幕与环境特征离散化**：自动生成符合市场规律的分辨率、色彩深度（如 1920-1080-1040-24）与可用显示区参数。",
                "**Client Hints 动态对齐**：精准匹配 Sec-CH-UA、User-Agent 与操作系统平台，避免出现「Linux 内核搭配 Windows UA」的硬伤。"
            ]
        },
        {
            "title": "📊 监控层：D3 指纹偏移与本地缓存持久化",
            "accent": C_EMERALD,
            "bullets": [
                "**D3.js 指纹偏移可视化**：前端直观呈现各批次任务的指纹熵值与偏移波动，辅助运营者实时评估指纹离散度与健康状态。",
                "**磁盘级指纹缓存预热**：通过 identity_cache 模块管理本地指纹缓存，减少重复运算开销，保证单账号注册期间环境绝对一致。",
                "**动态 AppJS 配置预热**：并发执行期间提前异步 WarmAppJSConfig，大幅减少注册前置握手时延。"
            ]
        }
    ]

    for i, ac in enumerate(anti_cards):
        ax = Inches(0.8 + i * 3.98)
        ay = Inches(2.05)
        aw = Inches(3.78)
        ah = Inches(4.8)

        c, ctf = add_card(s6, ax, ay, aw, ah, ac["title"], ac["accent"])
        p0 = ctf.paragraphs[0]
        p0.font.size = Pt(13)
        p0.space_after = Pt(12)

        for b in ac["bullets"]:
            p = ctf.add_paragraph()
            p.space_after = Pt(10)
            parts = b.split("**", 2)
            if len(parts) >= 3:
                r1 = p.add_run()
                r1.text = "• " + parts[1]
                r1.font.name = FONT_TITLE
                r1.font.size = Pt(11)
                r1.font.bold = True
                r1.font.color.rgb = C_TEXT_WHITE

                r2 = p.add_run()
                r2.text = parts[2]
                r2.font.name = FONT_BODY
                r2.font.size = Pt(10)
                r2.font.color.rgb = C_TEXT_MUTED
            else:
                r = p.add_run()
                r.text = "• " + b
                r.font.name = FONT_BODY
                r.font.size = Pt(10)
                r.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 7: MULTI-CHANNEL EMAIL MATRIX
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "邮箱生态 · EMAIL ECOSYSTEM", "高可用的全渠道邮箱矩阵与验证码自动化提取", "解耦的统一收件接口规范，无缝对接主流商业邮箱、临时邮件与私有自建服务")

    emails = [
        {
            "name": "🏢 Outlook 邮箱池",
            "mode": "IMAP / Graph REST API 双模式",
            "accent": C_INDIGO,
            "desc": "支持标准 IMAP 协议与现代 Microsoft Graph 接口。支持 RefreshToken 自动轮换刷新，从邮件正文精准正则提取 6 位验证码，适合高权重独享账号池。"
        },
        {
            "name": "☁️ iCloud 消息队列",
            "mode": "邮箱----消息列表URL 模式",
            "accent": C_CYAN,
            "desc": "创新集成专用消息列表 URL 取码通道，无需网页频繁输入双重认证密码，通过兼容的 API 端点秒级检索验证邮件，极其稳定高效。"
        },
        {
            "name": "📮 MoeMail 临时邮箱",
            "mode": "RESTful API 自动化集成",
            "accent": C_EMERALD,
            "desc": "对接 MoeMail 临时邮件服务，支持多域名配置与轮换机制；提供随机域名、全部域名广播与指定域名模式，有效避免单个后缀被风控靶向封锁。"
        },
        {
            "name": "🛠️ Cloud-Mail 自部署系统",
            "mode": "开源邮件服务器无缝接入",
            "accent": C_AMBER,
            "desc": "深度集成开源 cloud-mail 项目，支持通过管理员凭据自动从服务器拉取可用域名列表；私有化掌控度高，可无限扩展企业自定义域名后缀。"
        },
        {
            "name": "🏪 MailNest 商业邮箱",
            "mode": "商业级 Outlook 临时邮箱",
            "accent": C_ROSE,
            "desc": "对接 MailNest-迈巢商业邮件平台，界面直显账户余额；支持按项目代码动态采购并分配 Outlook 临时邮箱，免除自建维护负担。"
        }
    ]

    for i, em in enumerate(emails):
        ey = Inches(2.05 + i * 0.98)
        c, ctf = add_card(s7, Inches(0.8), ey, Inches(11.7), Inches(0.9), "", C_TEXT_WHITE, C_CARD_BG, C_CARD_BORDER)

        p0 = ctf.paragraphs[0]
        r_name = p0.add_run()
        r_name.text = em["name"] + "   "
        r_name.font.name = FONT_TITLE
        r_name.font.size = Pt(13)
        r_name.font.bold = True
        r_name.font.color.rgb = em["accent"]

        r_mode = p0.add_run()
        r_mode.text = "[ " + em["mode"] + " ]"
        r_mode.font.name = FONT_BODY
        r_mode.font.size = Pt(11)
        r_mode.font.color.rgb = C_CYAN
        p0.space_after = Pt(3)

        p1 = ctf.add_paragraph()
        p1.text = em["desc"]
        p1.font.name = FONT_BODY
        p1.font.size = Pt(10)
        p1.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 8: PROXY MANAGEMENT & ROUTING
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "网络治理 · PROXY & ROUTING", "住宅代理池调度与网络纯净度治理", "解决 OTP 400 与网络风控的核心基石：多协议支持、出口体检与任务级隔离")

    proxy_cards = [
        {
            "title": "🔌 全协议支持与批量管理",
            "accent": C_INDIGO,
            "items": [
                "**协议全覆盖**：全面支持 SOCKS5、HTTP、HTTPS 代理协议，保障各种复杂内网与隧道穿透需求。",
                "**认证格式规范**：支持 `protocol://user:pass@host:port` 格式，支持单条录入与文本批量秒级解析。",
                "**一键启停管控**：直观的代理状态切换开关，随时隔离异常节点，保障活跃代理池的高可用性。"
            ]
        },
        {
            "title": "🌍 实时出口检测与网络体检",
            "accent": C_CYAN,
            "items": [
                "**真实出口探测**：自动请求全球出口检测端点，解析出公网真实 IP 与所在国家/地区/城市坐标。",
                "**运营商与 ASN 识别**：明确区分住宅宽带 (Residential) 与机房数据中心 (Datacenter)，指导网络选型。",
                "**毫秒级延迟测速**：直观展示 TCP 握手与 HTTP 往返延迟 (Ping/Latency)，智能标记最优链路。"
            ]
        },
        {
            "title": "🛡️ 任务隔离与风控排查体系",
            "accent": C_AMBER,
            "items": [
                "**任务级代理绑定**：创建批处理任务时支持指定代理通道或直连，不同批次任务环境严格物理隔离。",
                "**OTP 400 智能诊断**：针对常见的 OTP 400 拦截，精准提示更换高纯净度住宅代理，消除盲目试错。",
                "**特别鸣谢生态合作**：原生适配 IPWO 覆盖 195+ 国家与地区的住宅代理生态，打通出海网络生命线。"
            ]
        }
    ]

    for i, pc in enumerate(proxy_cards):
        px = Inches(0.8 + i * 3.98)
        py = Inches(2.05)
        pw = Inches(3.78)
        ph = Inches(4.8)

        c, ctf = add_card(s8, px, py, pw, ph, pc["title"], pc["accent"])
        p0 = ctf.paragraphs[0]
        p0.font.size = Pt(13)
        p0.space_after = Pt(12)

        for item in pc["items"]:
            p = ctf.add_paragraph()
            p.space_after = Pt(10)
            parts = item.split("**", 2)
            if len(parts) >= 3:
                r1 = p.add_run()
                r1.text = "• " + parts[1]
                r1.font.name = FONT_TITLE
                r1.font.size = Pt(11)
                r1.font.bold = True
                r1.font.color.rgb = C_TEXT_WHITE

                r2 = p.add_run()
                r2.text = parts[2]
                r2.font.name = FONT_BODY
                r2.font.size = Pt(10)
                r2.font.color.rgb = C_TEXT_MUTED
            else:
                r = p.add_run()
                r.text = "• " + item
                r.font.name = FONT_BODY
                r.font.size = Pt(10)
                r.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 9: TASK SCHEDULING & RESILIENCE
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "调度与容错 · RESILIENCE", "高并发任务编排与资产零损耗容错体系", "在对抗网络波动与平台风控的同时，实现任务调度的高吞吐与业务资产的极致安全")

    res_quads = [
        {
            "badge": "⚡ 弹性并发控制",
            "title": "Worker Pool 与节流延时",
            "color": C_INDIGO,
            "desc": "支持自由配置注册总数量、并发 Worker 数量以及串行批次间等待延时。既可榨干网络吞吐极限，又可通过节奏拟人化延时有效削峰，降低被上游风控识别为机器刷量的概率。"
        },
        {
            "badge": "🛑 优雅中断与安全释放",
            "title": "全链路 Context 级即时响应",
            "color": C_CYAN,
            "desc": "注册核心各步骤全面监听 Go context.Context。用户在界面点击「停止」时，系统毫秒级安全中断所有正在进行的网络请求，立刻回收协程与连接池资源，不留僵尸进程。"
        },
        {
            "badge": "🔐 密码状态原子标记",
            "title": "避免二次消耗的退避重试",
            "color": C_EMERALD,
            "desc": "密码一旦在 AWS 端设置成功，立刻打上 passwordSet: true 原子标记。后置步骤（SSO 或 Token 交换）若遭遇抖动，仅执行原地 3 次重试，绝不将邮箱重新置回初始态浪费。"
        },
        {
            "badge": "📂 数据域物理隔离",
            "title": "凭据安全落盘与日志敏感脱敏",
            "color": C_AMBER,
            "desc": "应用运行时严格分离业务数据 (data/)、浏览器指纹缓存 (cache/) 与注册成果 (accounts.json)；运行日志全链路正则过滤敏感 API 终结点与 Token 凭证，兼具易调试性与安全性。"
        }
    ]

    for i, rq in enumerate(res_quads):
        row = i // 2
        col = i % 2
        rx = Inches(0.8 + col * 5.95)
        ry = Inches(2.05 + row * 2.45)
        rw = Inches(5.75)
        rh = Inches(2.25)

        c, ctf = add_card(s9, rx, ry, rw, rh, rq["badge"] + "  |  " + rq["title"], rq["color"])
        p0 = ctf.paragraphs[0]
        p0.font.size = Pt(13.5)
        p0.font.bold = True
        p0.space_after = Pt(8)

        p = ctf.add_paragraph()
        p.text = rq["desc"]
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 10: DESKTOP UI & USER EXPERIENCE
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)
    add_header(s10, "产品体验 · USER EXPERIENCE", "现代化跨平台桌面交互与运营级易用性", "摒弃传统脚本黑盒操控，提供兼具设计美学与运营监控能力的专业桌面环境")

    ui_features = [
        {
            "title": "📊 六大核心业务看板",
            "accent": C_INDIGO,
            "desc": "• 【概览看板】：环形任务进度图、即时成功率、实时统计、一键启停操作台\n• 【运行日志】：实时多级日志滚屏输出，支持格式化过滤与高亮提示\n• 【邮箱池】：Outlook、iCloud、MoeMail、CloudMail、MailNest 分标签统一管理\n• 【IP 管理】：代理节点快速增删改查、出口归属地与延迟一键测速\n• 【系统设置】：自定义数据存储根目录、成果输出路径与默认网络策略\n• 【关于页面】：自动检测 GitHub Releases 语义化版本更新"
        },
        {
            "title": "🎨 双色主题与极客体验",
            "accent": C_CYAN,
            "desc": "• 深色极客暗黑主题：深邃深蓝底色搭配霓虹高光，符合技术人员视觉习惯\n• 浅色清爽明亮主题：柔和 Slate 配色，高对比度排版，适应各种光照办公环境\n• 超轻体积冷启动：打包体积仅 20MB+，启动毫秒级响应，告别 Electron 臃肿体验"
        },
        {
            "title": "🌐 无缝三语国际化 (i18n)",
            "accent": C_EMERALD,
            "desc": "• 原生内置 简体中文 (zh-CN)、English (en-US)、日本語 (ja-JP)\n• 单页应用内无刷新即时平滑切换，海内外用户与开发者开箱即用"
        },
        {
            "title": "🔔 沉浸式完成提醒与通知",
            "accent": C_AMBER,
            "desc": "• 批量注册任务完成自动触发系统桌面级通知与专属提示音\n• 告别枯燥盯盘等待，打造真正可后台挂机静默运行的生产力工具"
        }
    ]

    for i, uf in enumerate(ui_features):
        row = i // 2
        col = i % 2
        ux = Inches(0.8 + col * 5.95)
        uy = Inches(2.05 + row * 2.45)
        uw = Inches(5.75)
        uh = Inches(2.25)

        c, ctf = add_card(s10, ux, uy, uw, uh, uf["title"], uf["accent"])
        p0 = ctf.paragraphs[0]
        p0.font.size = Pt(13.5)
        p0.font.bold = True
        p0.space_after = Pt(8)

        p = ctf.add_paragraph()
        p.text = uf["desc"]
        p.font.name = FONT_BODY
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 11: ROADMAP & 2API VISION
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11)
    add_header(s11, "战略演进 · ROADMAP & VISION", "未来规划：从桌面工具向企业级 2API 平台进化", "打破桌面客户端边界，将 Kiro 与 CodeWhisperer 能力深度融入现代 AI 编程全生态")

    roadmaps = [
        {
            "stage": "Milestone 01",
            "title": "桌面 GUI 走向无头 WebUI",
            "accent": C_INDIGO,
            "desc": "• 架构彻底前后端解耦，提供独立 Web 服务与 RESTful 控制接口\n• 支持 Docker 容器化一键部署与云端无头 (Headless) 自动化挂机\n• 多租户权限隔离与远程分布式集中注册池管理"
        },
        {
            "stage": "Milestone 02 · 核心重磅",
            "title": "内置 2API 服务网关",
            "accent": C_CYAN,
            "desc": "• 将 AWS CodeWhisperer / Kiro 能力转换为标准 OpenAI API 格式\n• 同步兼容 Anthropic Claude API 协议端点规范\n• 零门槛赋能 Cursor、Cline、Continue、VS Code 等主流 AI 编程套件"
        },
        {
            "stage": "Milestone 03",
            "title": "凭据全生命周期自动化调度",
            "accent": C_EMERALD,
            "desc": "• 自动定时无感刷新 Refresh Token 凭据，持续追踪账号存活\n• 智能多账号负载均衡池（支持轮询、最小剩余额度优先算法）\n• 异常凭据自动熔断与新号自动替补触发机制"
        },
        {
            "stage": "Milestone 04",
            "title": "自适应风控智能对抗",
            "accent": C_AMBER,
            "desc": "• 持续迭代协议层特征库，紧随云厂商风控规则演进\n• 结合多模态识别与智能对抗模型，自适应穿透新型 CAPTCHA 与质询\n• 打造高可靠、高弹性的出海与 AI 基础设施解决方案"
        }
    ]

    for i, rm in enumerate(roadmaps):
        rx = Inches(0.8 + i * 2.98)
        ry = Inches(2.05)
        rw = Inches(2.78)
        rh = Inches(4.8)

        c, ctf = add_card(s11, rx, ry, rw, rh, rm["stage"], rm["accent"])
        p0 = ctf.paragraphs[0]
        p0.font.size = Pt(12)
        p0.space_after = Pt(6)

        p1 = ctf.add_paragraph()
        p1.text = rm["title"]
        p1.font.name = FONT_TITLE
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = C_TEXT_WHITE
        p1.space_after = Pt(14)

        p2 = ctf.add_paragraph()
        p2.text = rm["desc"]
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 12: SUMMARY & COMMUNITY
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12)
    
    # Big decorative glow card
    center_card = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(0.8), Inches(10.933), Inches(5.9))
    center_card.fill.solid()
    center_card.fill.fore_color.rgb = C_CARD_BG
    center_card.line.color.rgb = C_INDIGO
    center_card.line.width = Pt(1.5)

    if os.path.exists(icon_path):
        s12.shapes.add_picture(icon_path, Inches(2.0), Inches(1.5), width=Inches(1.8), height=Inches(1.8))

    c_tf = center_card.text_frame
    c_tf.word_wrap = True
    c_tf.margin_left = Inches(3.0)
    c_tf.margin_top = Inches(0.8)

    p = c_tf.paragraphs[0]
    p.text = "KiroX"
    p.font.name = FONT_TITLE
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_WHITE

    p2 = c_tf.add_paragraph()
    p2.text = "定义新一代极速、健壮、开放的 AI 凭据基础设施"
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = C_CYAN
    p2.space_before = Pt(4)
    p2.space_after = Pt(18)

    p3 = c_tf.add_paragraph()
    p3.text = "• 卓越性能：纯 Go 打造，超低资源占用，毫秒级网络吞吐响应\n• 深度工程：完整还原 JWE 非对称加密、OIDC 与 SSO 核心认证流\n• 立体防护：领先的 TLS/JA3/HTTP2 指纹仿真与浏览器时延分布建模\n• 开源共享：基于 Apache License 2.0 协议，欢迎交流共建"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(12)
    p3.font.color.rgb = C_TEXT_HIGHLIGHT
    p3.space_after = Pt(20)

    # 3 Pill buttons at bottom
    p_links = [
        ("GitHub 仓库", "github.com/huey1in/kirox", C_INDIGO),
        ("开源协议", "Apache License 2.0", C_CYAN),
        ("交流与共建", "QQ 交流群 / 欢迎提交 Issue & PR", C_EMERALD)
    ]
    for j, (l_title, l_desc, l_col) in enumerate(p_links):
        bx = Inches(2.0 + j * 3.2)
        by = Inches(5.1)
        bshape = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, Inches(3.0), Inches(0.8))
        bshape.fill.solid()
        bshape.fill.fore_color.rgb = C_BG
        bshape.line.color.rgb = l_col
        bshape.line.width = Pt(1)

        btf = bshape.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp0 = btf.paragraphs[0]
        bp0.text = l_title
        bp0.font.name = FONT_TITLE
        bp0.font.size = Pt(12)
        bp0.font.bold = True
        bp0.font.color.rgb = l_col
        bp0.alignment = PP_ALIGN.CENTER

        bp1 = btf.add_paragraph()
        bp1.text = l_desc
        bp1.font.name = FONT_BODY
        bp1.font.size = Pt(10)
        bp1.font.color.rgb = C_TEXT_MUTED
        bp1.alignment = PP_ALIGN.CENTER

    output_path = "KiroX_项目介绍.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_deck()
